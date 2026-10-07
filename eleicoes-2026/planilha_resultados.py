#!/usr/bin/env python3
"""Gera resultados-2026.xlsx a partir dos CSVs de consolidar_tse.py.

Abas:
  Leia-me
  Resumo Estados / Resumo Municípios   aptos, comparecimento, abstenção,
                                       brancos, nulos e válidos por cargo
  Presidente / Governador / Senador    votos de cada candidato por município
  Presidente (UF) ... Dep. Distrital   votos de cada candidato por estado
  Comparação Estados / Comparação Municípios
                                       Presidente x Governador

Uso:
  python3 baixar_tse.py --dir ./tse-json
  python3 consolidar_tse.py --dir ./tse-json --saida ./dados
  python3 planilha_resultados.py --dados ./dados
"""
import argparse
import csv
import os
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

AQUI = os.path.dirname(os.path.abspath(__file__))
SAIDA = os.path.join(AQUI, "resultados-2026.xlsx")
FONTE = ("Fonte: TSE, resultados.tse.jus.br (eleições 6257 e 6259, 1º turno de 04/10/2026), "
         "arquivos de totalização de 05/10/2026.")
ORDEM_REG = ["Norte", "Nordeste", "Centro-Oeste", "Sudeste", "Sul", "Exterior"]
CARGOS_UF = ["Presidente", "Governador", "Senador", "Deputado Federal", "Deputado Estadual",
             "Deputado Distrital"]

FONT = "Arial"
F_BASE = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name=FONT, size=14, bold=True)
F_NOTE = Font(name=FONT, size=9, italic=True, color="555555")
F_LINK = Font(name=FONT, size=10, color="008000")
FILL_HEAD = PatternFill("solid", fgColor="1F3A5F")
NUM = '#,##0;-#,##0;"-"'
PCT = '0.00%;-0.00%;"-"'


def ler(dados, nome):
    with open(os.path.join(dados, f"{nome}.csv"), encoding="utf-8") as f:
        return list(csv.DictReader(f, delimiter=";"))


def chave_ordem(r):
    reg = ORDEM_REG.index(r["regiao"]) if r["regiao"] in ORDEM_REG else 99
    return (reg, r["uf"], r["municipio"], r["cargo"])


class Aba:
    """Aba em modo write_only com cabeçalho, larguras e formatos por coluna."""

    def __init__(self, wb, titulo, colunas, nota=None):
        self.ws = wb.create_sheet(titulo)
        self.fmts = []
        self.linha = 0
        for i, (_nome, larg, fmt) in enumerate(colunas, start=1):
            self.ws.column_dimensions[get_column_letter(i)].width = larg
            self.fmts.append(fmt)
        self.ws.freeze_panes = "A3" if nota else "A2"
        if nota:
            self._raw([self._cel(nota, F_NOTE)])
        cab = []
        for nome, _l, _f in colunas:
            c = self._cel(nome, F_HEAD)
            c.fill = FILL_HEAD
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cab.append(c)
        self._raw(cab)
        self.primeira = self.linha + 1

    def _cel(self, valor, fonte=F_BASE, fmt=None):
        c = WriteOnlyCell(self.ws, value=valor)
        c.font = fonte
        if fmt:
            c.number_format = fmt
        return c

    def _raw(self, cels):
        self.ws.append(cels)
        self.linha += 1

    def add(self, valores, fonte=F_BASE, fontes=None):
        self.linha += 1
        r = self.linha
        cels = []
        for i, v in enumerate(valores):
            if isinstance(v, str) and "{r}" in v:
                v = v.replace("{r}", str(r))
            f = (fontes or {}).get(i, fonte)
            cels.append(self._cel(v, f, self.fmts[i] if i < len(self.fmts) else None))
        self.ws.append(cels)


def aba_leiame(wb, totais):
    ws = wb.create_sheet("Leia-me")
    ws.column_dimensions["A"].width = 120
    linhas = [
        ("Eleições Gerais 2026 · 1º turno (04/10/2026) · resultados oficiais do TSE", F_TITLE),
        ("", F_BASE),
        (f"Presidente, Brasil (com exterior): Flávio Bolsonaro {totais['FLAVIO BOLSONARO']:,} votos; "
         f"Lula {totais['LULA']:,} votos.".replace(",", "."), F_BOLD),
        ("", F_BASE),
        ("Abas", Font(name=FONT, size=11, bold=True)),
        ("Resumo Estados / Resumo Municípios: aptos, comparecimento, abstenção, brancos, nulos e válidos "
         "por cargo, com percentuais calculados.", F_BASE),
        ("Presidente, Governador, Senador: votos de cada candidato, brancos e nulos em cada município. "
         "Use o filtro do cabeçalho para escolher estado ou cidade.", F_BASE),
        ("Abas '(UF)' e Deputados: votos por estado. Deputados vêm só por estado: por município seriam "
         "milhões de linhas, acima do limite do Excel.", F_BASE),
        ("Comparação Estados / Comparação Municípios: Presidente x Governador, com a diferença e se bate.",
         F_BASE),
        ("", F_BASE),
        ("Como ler a comparação", Font(name=FONT, size=11, bold=True)),
        ("O comparecimento para Presidente costuma ser um pouco maior que para Governador: quem vota em "
         "trânsito fora do seu estado só pode votar para Presidente (o TSE avisa isso na tela do boletim "
         "de urna). Brancos, nulos e válidos variam porque o eleitor pode anular em um cargo e não no "
         "outro. No Senador há 2 vagas, então cada eleitor dá 2 votos.", F_BASE),
        ("", F_BASE),
        ("Legenda de cores", Font(name=FONT, size=11, bold=True)),
        ("Texto preto: dados do TSE ou fórmulas. Texto verde: fórmula que busca dado em outra aba.", F_BASE),
        ("", F_BASE),
        (FONTE, F_NOTE),
        ("Para atualizar: python3 eleicoes-2026/baixar_tse.py, consolidar_tse.py e planilha_resultados.py "
         "(veja o cabeçalho de cada script).", F_NOTE),
    ]
    for t, f in linhas:
        c = WriteOnlyCell(ws, value=t)
        c.font = f
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.append([c])


COLS_RESUMO = [("Região", 13, None), ("UF", 6, None), ("Cód. TSE", 10, None), ("Local", 28, None),
               ("Cargo", 18, None), ("Aptos", 14, NUM), ("Comparecimento", 15, NUM),
               ("Abstenção", 14, NUM), ("Brancos", 13, NUM), ("Nulos", 13, NUM),
               ("Válidos", 14, NUM), ("Votos de legenda", 13, NUM), ("% Abstenção", 11, PCT),
               ("% Brancos", 10, PCT), ("% Nulos", 10, PCT)]


def aba_resumo(wb, titulo, linhas, nota):
    a = Aba(wb, titulo, COLS_RESUMO, nota)
    for r in sorted(linhas, key=chave_ordem):
        a.add([r["regiao"], r["uf"], r["cod_municipio"], r["municipio"], r["cargo"],
               int(r["aptos"]), int(r["comparecimento"]), int(r["abstencao"]), int(r["brancos"]),
               int(r["nulos"]), int(r["validos"]), int(r["legenda"]),
               '=IF(F{r}=0,"",H{r}/F{r})', '=IF(G{r}=0,"",I{r}/G{r})', '=IF(G{r}=0,"",J{r}/G{r})'])
    a.ws.auto_filter.ref = f"A2:O{a.linha}"
    return a


COLS_VOTOS = [("Região", 13, None), ("UF", 6, None), ("Cód. TSE", 10, None), ("Local", 28, None),
              ("Tipo", 11, None), ("Número", 9, None), ("Candidato", 30, None), ("Partido", 14, None),
              ("Situação", 16, None), ("Votos", 14, NUM)]


def aba_votos(wb, titulo, linhas, nota):
    a = Aba(wb, titulo, COLS_VOTOS, nota)
    ordem_tipo = {"Candidato": 0, "Legenda": 1, "Branco": 2, "Nulo": 3}
    linhas = sorted(linhas, key=lambda r: (chave_ordem(r), ordem_tipo.get(r["tipo"], 9), -int(r["votos"])))
    for r in linhas:
        a.add([r["regiao"], r["uf"], r["cod_municipio"], r["municipio"], r["tipo"], r["numero"],
               r["candidato"], r["partido"], r["situacao"], int(r["votos"])])
    a.ws.auto_filter.ref = f"A2:J{a.linha}"


def aba_comp_estados(wb, res_uf, res_mun_linhas):
    cols = [("Região", 13, None), ("UF", 6, None), ("Estado", 20, None),
            ("Comparecimento Presidente", 15, NUM), ("Comparecimento Governador", 15, NUM),
            ("Diferença", 12, NUM), ("Diferença %", 10, PCT), ("Comparecimento igual?", 14, None),
            ("Válidos Presidente", 15, NUM), ("Válidos Governador", 15, NUM), ("Diferença", 12, NUM),
            ("Brancos+Nulos Presidente", 15, NUM), ("Brancos+Nulos Governador", 15, NUM),
            ("Diferença", 12, NUM), ("Soma municípios Presidente (válidos)", 16, NUM),
            ("Bate com estado?", 12, None), ("Soma municípios Governador (válidos)", 16, NUM),
            ("Bate com estado?", 12, None)]
    a = Aba(wb, "Comparação Estados", cols,
            "Presidente x Governador por estado. Diferença = Governador - Presidente. Fórmulas buscam "
            "nas abas Resumo Estados e Resumo Municípios.")
    rs, rm = "'Resumo Estados'", "'Resumo Municípios'"
    ufs = sorted({(r["regiao"], r["uf"], r["municipio"]) for r in res_uf if r["uf"] != "ZZ"},
                 key=lambda x: (ORDEM_REG.index(x[0]), x[1]))

    def sif(aba, col, uf, cargo):
        return f'SUMIFS({aba}!${col}:${col},{aba}!$B:$B,"{uf}",{aba}!$E:$E,"{cargo}")'

    for reg, uf, nome in ufs:
        P, G = "Presidente", "Governador"
        a.add([reg, uf, nome,
               "=" + sif(rs, "G", uf, P), "=" + sif(rs, "G", uf, G),
               "=E{r}-D{r}", '=IF(D{r}=0,"",F{r}/D{r})',
               '=IF(F{r}=0,"Igual","Diferente")',
               "=" + sif(rs, "K", uf, P), "=" + sif(rs, "K", uf, G), "=J{r}-I{r}",
               f"={sif(rs, 'I', uf, P)}+{sif(rs, 'J', uf, P)}",
               f"={sif(rs, 'I', uf, G)}+{sif(rs, 'J', uf, G)}", "=M{r}-L{r}",
               "=" + sif(rm, "K", uf, P), '=IF(O{r}=I{r},"OK","Não bate")',
               "=" + sif(rm, "K", uf, G), '=IF(Q{r}=J{r},"OK","Não bate")'],
              fontes={i: F_LINK for i in (3, 4, 8, 9, 11, 12, 14, 16)})
    fim = a.linha
    tot = ["Total", "", "27 UFs"]
    for i in range(3, 18):
        L = get_column_letter(i + 1)
        tot.append(f"=SUM({L}3:{L}{fim})" if i in (3, 4, 5, 8, 9, 10, 11, 12, 13, 14, 16) else
                   ('=IF(D{r}=0,"",F{r}/D{r})' if i == 6 else ""))
    a.add(tot, fonte=F_BOLD)
    a.add(["Estados com comparecimento diferente:", "", "", "", "", "", "",
           f'=COUNTIF(H3:H{fim},"Diferente")'], fonte=F_BOLD)


def aba_comp_municipios(wb, res_mun):
    cols = [("Região", 13, None), ("UF", 6, None), ("Cód. TSE", 10, None), ("Município", 28, None),
            ("Comparecimento Presidente", 15, NUM), ("Comparecimento Governador", 15, NUM),
            ("Diferença", 11, NUM), ("Comparecimento igual?", 13, None),
            ("Válidos Presidente", 14, NUM), ("Válidos Governador", 14, NUM), ("Diferença", 11, NUM),
            ("Brancos+Nulos Presidente", 14, NUM), ("Brancos+Nulos Governador", 14, NUM),
            ("Diferença", 11, NUM)]
    a = Aba(wb, "Comparação Municípios", cols,
            "Presidente x Governador por município (números do TSE; diferenças calculadas por fórmula). "
            "Diferença no comparecimento costuma ser voto em trânsito.")
    por = defaultdict(dict)
    for r in res_mun:
        por[(r["regiao"], r["uf"], r["cod_municipio"], r["municipio"])][r["cargo"]] = r
    for k in sorted(por, key=lambda k: (ORDEM_REG.index(k[0]), k[1], k[3])):
        p, g = por[k].get("Presidente"), por[k].get("Governador")
        if not p or not g:
            continue
        a.add(list(k) + [int(p["comparecimento"]), int(g["comparecimento"]), "=F{r}-E{r}",
                         '=IF(G{r}=0,"Igual","Diferente")',
                         int(p["validos"]), int(g["validos"]), "=J{r}-I{r}",
                         int(p["brancos"]) + int(p["nulos"]), int(g["brancos"]) + int(g["nulos"]),
                         "=M{r}-L{r}"])
    a.ws.auto_filter.ref = f"A2:N{a.linha}"


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--dados", required=True, help="pasta com os CSVs de consolidar_tse.py")
    p.add_argument("--saida", default=SAIDA)
    p.add_argument("--ufs", help="só estas UFs (ex.: SP,AC), para testes")
    a = p.parse_args()
    filtro = set(a.ufs.upper().split(",")) if a.ufs else None

    def ok(r):
        return not filtro or r["uf"] in filtro

    vm = [r for r in ler(a.dados, "votos_municipio") if ok(r)]
    ve = [r for r in ler(a.dados, "votos_estado") if ok(r)]
    rm = [r for r in ler(a.dados, "resumo_municipio") if ok(r)]
    re_ = [r for r in ler(a.dados, "resumo_estado") if ok(r)]
    totais = defaultdict(int)
    for r in ve:
        if r["cargo"] == "Presidente":
            totais[r["candidato"]] += int(r["votos"])

    wb = Workbook(write_only=True)
    aba_leiame(wb, totais)
    aba_resumo(wb, "Resumo Estados", re_, FONTE + " Exterior (ZZ) só tem Presidente.")
    aba_resumo(wb, "Resumo Municípios", rm, FONTE)
    for cargo in ("Presidente", "Governador", "Senador"):
        aba_votos(wb, cargo, [r for r in vm if r["cargo"] == cargo],
                  f"{cargo}: votos por município. " + FONTE)
    nomes_uf = {"Presidente": "Presidente (UF)", "Governador": "Governador (UF)", "Senador": "Senador (UF)",
                "Deputado Federal": "Dep. Federal (UF)", "Deputado Estadual": "Dep. Estadual (UF)",
                "Deputado Distrital": "Dep. Distrital (DF)"}
    for cargo in CARGOS_UF:
        linhas = [r for r in ve if r["cargo"] == cargo]
        if linhas:
            aba_votos(wb, nomes_uf[cargo], linhas, f"{cargo}: votos por estado. " + FONTE)
    aba_comp_estados(wb, re_, rm)
    aba_comp_municipios(wb, rm)
    wb.save(a.saida)
    print(f"Gerado: {a.saida}")


if __name__ == "__main__":
    main()
